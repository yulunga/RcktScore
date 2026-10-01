from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/glennrowe/Development/Projects/RcktScore")
OUTPUT = ROOT / "testing/reports/HitnScore-Production-Release-Readiness-and-Four-Day-Plan-2026-09-11.docx"
BRAND_LOGO = ROOT / "frontend/public/branding/logo/brand-logo.png"
PRIOR_REPORT = ROOT / "testing/reports/RcktScore-Apple-Subscription-Sandbox-Test-Report-2026-09-11.docx"

BLACK = "000000"
INK = "102A43"
NAVY = "0B4F88"
BLUE = "2F80ED"
PINK = "EC5EA8"
MID_GREY = "486581"
LIGHT_GREY = "D9E2EC"
PALE_BLUE = "F4F8FC"
PALE_GREY = "F7F9FB"
GREEN = "138A4A"
AMBER = "9A6700"
RED = "C62828"


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def border(cell, color: str = LIGHT_GREY, size: str = "5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def margins(cell, top=90, start=105, bottom=90, end=105) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def page_number(paragraph) -> None:
    paragraph.add_run("Page ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, end])


def bottom_rule(paragraph, color: str = PINK) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def wordmark(paragraph, size: float = 18, suffix: str = "") -> None:
    for text, color in (("Hit", BLUE), ("n", PINK), ("Score", BLUE)):
        run = paragraph.add_run(text)
        run.bold = True
        run.font.name = "Aptos Display"
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(color)
    if suffix:
        run = paragraph.add_run(suffix)
        run.font.name = "Aptos"
        run.font.size = Pt(max(size - 4, 8))
        run.font.color.rgb = rgb(MID_GREY)


def table(doc: Document, headers: list[str], rows: Iterable[Iterable[str]], widths=None, font_size=8.7):
    rows = [list(row) for row in rows]
    result = doc.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    repeat_header(result.rows[0])
    for index, header in enumerate(headers):
        cell = result.rows[0].cells[index]
        cell.text = header
        shade(cell, NAVY)
        border(cell)
        margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = rgb("FFFFFF")
            run.font.size = Pt(font_size)
    for row_index, values in enumerate(rows):
        row = result.add_row()
        cant_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            border(cell)
            margins(cell)
            shade(cell, "FFFFFF" if row_index % 2 == 0 else PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    value_text = str(value).lower()
                    if value_text in {"pass", "passed", "complete", "closed"}:
                        run.bold = True
                        run.font.color.rgb = rgb(GREEN)
                    elif value_text in {"blocker", "no-go", "failed"}:
                        run.bold = True
                        run.font.color.rgb = rgb(RED)
                    elif value_text in {"conditional", "open", "required"}:
                        run.bold = True
                        run.font.color.rgb = rgb(AMBER)
        if widths:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    if widths:
        for index, width in enumerate(widths):
            result.rows[0].cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return result


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(item)


def status_line(doc: Document, label: str, text: str, color: str) -> None:
    paragraph = doc.add_paragraph()
    lead = paragraph.add_run(label + ": ")
    lead.bold = True
    lead.font.color.rgb = rgb(color)
    paragraph.add_run(text)


def configure(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = rgb(BLACK)

    for name, size in (("Heading 1", 18), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(BLACK)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name == "Heading 1" else 7)
        style.paragraph_format.space_after = Pt(4)

    for name in ("List Bullet", "List Number"):
        styles[name].font.name = "Aptos"
        styles[name].font.size = Pt(10.1)
        styles[name].paragraph_format.space_after = Pt(3)


def headers_and_footers(doc: Document) -> None:
    for section in doc.sections:
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        wordmark(header, 10.5, "  |  Production readiness · decision plan")
        bottom_rule(header)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("HitnScore  •  Production release readiness  •  ")
        run.font.size = Pt(8)
        run.font.color.rgb = rgb(MID_GREY)
        page_number(footer)
        for footer_run in footer.runs:
            footer_run.font.size = Pt(8)
            footer_run.font.color.rgb = rgb(MID_GREY)


def add_page(doc: Document, heading: str) -> None:
    doc.add_page_break()
    doc.add_heading(heading, level=1)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relationship = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def main() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.66)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    configure(doc)

    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand.add_run().add_picture(str(BRAND_LOGO), width=Inches(0.85))
    brand.paragraph_format.space_after = Pt(1)
    brand_name = doc.add_paragraph()
    brand_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    wordmark(brand_name, 24)
    brand_name.paragraph_format.space_after = Pt(13)

    title = doc.add_paragraph("Production Release Readiness Review and Four Day Plan", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Current-state comparison, controlled freeze, launch gates and Root Admin security")
    subtitle_run.bold = True
    subtitle_run.font.size = Pt(13)
    subtitle_run.font.color.rgb = rgb(BLACK)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Assessment date: ").bold = True
    meta.add_run("11 September 2026\n")
    meta.add_run("Target: ").bold = True
    meta.add_run("submission-ready build within four days; public availability subject to Apple review\n")
    meta.add_run("Release posture: ").bold = True
    meta.add_run("conditional go with a narrow launch scope and a backend/database freeze after the launch blockers are closed\n")
    meta.add_run("Evidence base: ").bold = True
    meta.add_run("repository state, recent commit history, checked-in readiness documents, subscription Test A–D report, and verification run on 11 September 2026")

    doc.add_page_break()
    doc.add_heading("Executive decision", level=1)
    status_line(doc, "Recommendation", "Proceed toward a controlled submission, but do not declare the product production-ready or freeze the backend today.", AMBER)
    doc.add_paragraph(
        "The product has moved materially closer to release since the 8 September readiness review. Apple purchase verification and lifecycle processing now exist, the manual Sandbox programme has proved Tests A through D, root-admin access uses expiring opaque server-side sessions, and core web/backend compilation checks pass. The remaining risk is concentrated in release operations, repeatable automated gates, production Apple configuration, untested exceptional subscription paths, native archive validation, and privileged-administration security."
    )
    doc.add_paragraph(
        "A four-day target is credible only if it means producing and submitting a deliberately narrow release candidate. Apple controls review duration and public availability, so no internal plan can guarantee that the app is live in exactly four days. The launch should be stopped if any P0 gate in this document remains open at the freeze checkpoint."
    )
    table(doc, ["Decision", "Position"], [
        ["Product scope", "Freeze to the implemented v1 experience: authentication, personal and club membership flows, squash/racketball/tennis, match lifecycle, Personal Free/Plus, settings, help and current notification inbox."],
        ["Backend/schema freeze", "Conditional. Close the P0 security, Apple configuration, build, data and monitoring gates first; then permit only explicitly approved launch-blocker fixes."],
        ["Style-only period", "Yes after the freeze, provided changes do not alter API contracts, entitlements, authentication, persistence, scoring rules or release configuration."],
        ["Root Admin", "Do not expose broadly in its current form. Add strong MFA-backed edge authorization or keep it network-restricted until that control is live."],
        ["Four-day outcome", "Target a signed release candidate and App Store submission. Public release on 15 September is conditional on review completion and a final go/no-go decision."],
    ], widths=[1.35, 5.55], font_size=9.2)

    doc.add_heading("Readiness estimate", level=1)
    doc.add_paragraph(
        "These percentages are planning estimates, not automated quality scores. They express how much of the narrow launch boundary is implemented and how much evidence exists to operate it safely."
    )
    table(doc, ["Measure", "8 Sep assessment", "11 Sep assessment", "Interpretation"], [
        ["Narrow v1 feature completeness", "75–80%", "85–90%", "Core product, native scoring and paid plan behaviour are substantially present; known scaffolds remain excluded."],
        ["Production readiness", "55–60%", "70–75%", "Subscription and admin foundations improved, but reproducible gates, production configuration, operations and privileged security still prevent unconditional sign-off."],
        ["Subscription implementation", "about 25%", "about 80%", "Verification, notifications, reconciliation, entitlements, admin activity and audit exist; exceptional cases and production activation remain."],
        ["Release confidence today", "Low", "Conditional", "There is enough product to submit, but not enough evidence to freeze or launch without the P0 plan."],
    ], widths=[1.55, 1.05, 1.1, 3.2])

    doc.add_heading("What changed since the previous review", level=1)
    table(doc, ["Date / milestone", "Material change", "Release effect"], [
        ["3 Sep · root-admin session work", "Opaque random root-admin bearer tokens are hashed in the database, expire, can be revoked, and replace the former trust-header bypass.", "Closes the most obvious authentication bypass, but does not add MFA, browser-safe cookies, throttling, idle expiry or privileged audit."],
        ["7–8 Sep · product/admin expansion", "Root-admin user, subscription, notification and sport-management functions expanded; native settings, scoring, offline replay and Personal Plus views matured.", "Feature completeness increased, while the privileged attack surface and regression surface also grew."],
        ["9–10 Sep · Apple lifecycle", "Signed transaction verification, App Store Server Notifications V2 handling, reconciliation, lifecycle state, entitlement audit and renewal-boundary safeguards were implemented.", "Paid entitlement is now server-authoritative and technically credible for release after production activation and remaining evidence."],
        ["10–11 Sep · manual Sandbox QA", "Tests A–D proved initial purchase, cancellation/expiry, renewal, failed renewal, grace, final expiry, account binding and billing recovery.", "Removes the main happy-path lifecycle uncertainty and identifies the smaller exceptional-path set still open."],
    ], widths=[1.2, 3.25, 2.45])
    doc.add_heading("Subscription programme result", level=2)
    table(doc, ["Test", "Scenario", "Result", "What it proves"], [
        ["A", "Monthly purchase, voluntary cancellation, paid access until expiry, downgrade", "Passed", "Initial entitlement, cancellation semantics and expiry downgrade."],
        ["B", "Yearly purchase and accelerated renewals", "Passed", "Stable renewals without plan churn."],
        ["C", "Failed renewal, grace period, grace expiry and billing-retry expiry", "Passed", "Grace access and terminal downgrade; account binding rejection also observed."],
        ["D", "Failed renewal followed by payment recovery", "Passed", "Grace-to-active recovery and subsequent renewal without entitlement audit noise."],
    ], widths=[0.45, 2.55, 0.8, 3.1])
    doc.add_paragraph(
        "The detailed screenshots, SQL evidence, rectifications and outcomes are retained in the companion Apple Subscription Sandbox Test Report dated 11 September 2026. That report signs off the manual suite performed; it does not sign off every production exception or configuration item."
    )

    doc.add_heading("Known exclusions from this release", level=2)
    bullets(doc, [
        "WebSocket live broadcast is scaffolded but not production-complete.",
        "APNs push delivery and background badge behaviour are not included; the notification inbox continues to poll.",
        "Padel, table tennis, badminton and pickleball engines remain fail-safe placeholders and must stay hidden/disabled.",
        "Federation-style association links, account-level game presets and reporting scaffolds are not release promises.",
        "Profile photos remain device-local and are not a cross-device feature.",
        "Offline capability remains limited to previously loaded active matches and queued scoring; new setup and unseen history still require connectivity.",
    ])

    doc.add_heading("Verification performed for this review", level=1)
    table(doc, ["Gate", "Result on 11 Sep", "Evidence / implication"], [
        ["Backend Python compilation", "Pass", "All backend/common and backend/functions Python files compiled successfully with an isolated bytecode cache."],
        ["Frontend production build", "Pass", "Vite production build completed: 69 modules; main JavaScript about 427.5 kB, 112.3 kB gzip."],
        ["Native tennis scenarios", "Pass", "The checked-in tennis scenario shell suite completed successfully."],
        ["Backend pytest", "Not executed", "The current environment lacks pytest. Historical pre-deployment evidence records 79 passing tests, but the release gate is not reproducible here today."],
        ["Playwright smoke", "Not executed", "The Playwright command is not installed in the current frontend environment."],
        ["Native Release build", "Environment blocked", "The generic iOS Release build reached target compilation but failed in asset catalog processing because CoreSimulator runtimes/services are unavailable; an Archive on a healthy Xcode host is mandatory."],
    ], widths=[1.45, 1.15, 4.3])
    status_line(doc, "Interpretation", "The code has useful positive signals, but there is not yet one reproducible green release pipeline. Missing tools and an unavailable simulator are release-process failures even when they are not source-code failures.", AMBER)

    doc.add_heading("Evidence still needed for paid production", level=2)
    table(doc, ["Evidence", "Minimum launch expectation", "Status"], [
        ["Refund, revoke and reversal", "Server notification changes entitlement correctly; reversal is applied only when Apple confirms it.", "Open"],
        ["Cross-device restore", "Same HitnScore account restores; another account cannot claim the transaction.", "Open"],
        ["Missed-notification repair", "Reconciliation repairs intentionally skipped lifecycle delivery without entitlement churn.", "Open"],
        ["Negative verification", "Invalid signature, bundle, app Apple ID, product, environment and account token are rejected and logged safely.", "Open"],
        ["Duplicate/out-of-order delivery", "Idempotent state remains correct under replay and reordering.", "Open"],
        ["Production configuration", "Migration 027, secret/key, Apple URLs, environment routing and grace policy verified against deployed resources.", "Open"],
        ["Monitoring and recovery", "Alarms, DLQ/replay path, runbook and responsible operator are exercised.", "Open"],
    ], widths=[1.55, 4.35, 0.9])

    doc.add_heading("The four-day release plan", level=1)
    doc.add_paragraph(
        "The schedule below starts on Friday 11 September 2026 and targets a submission-ready build by Monday 14 September, with controlled public release on Tuesday 15 September if Apple review and the final gate permit it. If the team starts later, preserve the sequence rather than compressing the gates."
    )
    table(doc, ["Day", "Primary objective", "Required work", "Exit criterion"], [
        ["Day 1 · Fri 11 Sep", "Fix the release boundary", "Name the release commit; assign owners; verify migrations 024–027; confirm backup/PITR; inventory production secrets and Apple configuration without exposing values; decide Root Admin containment; install and pin test dependencies.", "Scope signed, no new feature work, all P0 items owned, security architecture chosen."],
        ["Day 2 · Sat 12 Sep", "Close launch blockers", "Implement the Root Admin minimum; restrict CORS/security headers; run remaining subscription exception tests; make backend pytest and web smoke reproducible; configure alarms/DLQ/runbook; verify production Apple endpoints and release purchase flag.", "All backend/schema/security changes merged, deployed to staging, migrations verified, full automated and manual gate green."],
        ["Freeze · Sat 12 Sep 18:00", "Lock functions and data contracts", "Tag the backend/schema candidate; record migration checksums; snapshot configuration; stop ordinary Lambda, API, entitlement, scoring and schema changes.", "Named release tag, deploy manifest, rollback target and signed freeze decision exist."],
        ["Day 3 · Sun 13 Sep", "Validate the immutable candidate", "Create signed Xcode Archive; TestFlight install; test login, registration, deletion, membership selection, scoring, offline replay, history limits, purchase/restore and lifecycle; verify App Store metadata, privacy and review notes.", "No P0/P1 defects; release build and backend versions recorded; reviewer credentials and notes tested."],
        ["Day 4 · Mon 14 Sep", "Submit and rehearse operations", "Final backup; production deployment from the tagged artifact; smoke test; submit build/IAP; rehearse rollback, webhook replay and support response; begin heightened monitoring.", "Submission accepted by App Store Connect, production checks green, incident owner on call."],
        ["Target · Tue 15 Sep", "Controlled public release", "Review Apple result; re-run go/no-go; release manually or phased; monitor auth, purchases, notifications, errors and support.", "No open P0, error rates normal, purchase verification healthy, rollback immediately available."],
    ], widths=[1.1, 1.4, 3.45, 1.25], font_size=8.1)

    doc.add_page_break()
    doc.add_heading("Ownership model", level=2)
    table(doc, ["Role", "Accountability"], [
        ["Release owner", "Single go/no-go authority, scope control, freeze exceptions and App Store submission."],
        ["Backend/data owner", "Deployable artifact, migrations, Apple endpoints, reconciliation, monitoring, rollback and data recovery."],
        ["iOS owner", "Archive, device matrix, TestFlight verification, privacy manifest, StoreKit gate and review notes."],
        ["Web/admin owner", "Frontend build, Root Admin delivery/containment, CORS/CSP and privileged workflow tests."],
        ["QA recorder", "Evidence ledger, test accounts, screenshots, exact build/commit/environment and defect disposition."],
    ], widths=[1.35, 5.55])

    doc.add_heading("P0 launch gates", level=1)
    table(doc, ["Gate", "Pass condition", "Decision if not passed"], [
        ["Root Admin protection", "MFA-backed authorization is enforced at the API edge, or the portal/API are inaccessible from the public internet and available only through an approved restricted path.", "No-go"],
        ["Production Apple activation", "Agreements, banking/tax, products, key/secret, bundle/app IDs, notification URLs, grace policy and production environment routing are independently verified.", "No-go for paid launch"],
        ["Database and migration safety", "Migrations 024–027 applied once, backup/PITR confirmed, schema version recorded, rollback/forward-fix procedure rehearsed.", "No-go"],
        ["Subscription exceptions", "Refund/revoke, cross-device restore, missed-notification reconciliation, negative verification and duplicate/out-of-order tests pass.", "No-go for paid launch"],
        ["Repeatable automated gate", "Backend tests, frontend build/smoke and native logic tests run from documented commands on a clean release machine.", "No-go"],
        ["Signed iOS archive", "Release Archive succeeds, uploads, installs via TestFlight and passes the device/UAT checklist with the exact backend candidate.", "No-go"],
        ["Monitoring and incident readiness", "Structured logs, alarms, notification failure handling, replay/DLQ path, support contact and rollback owner are active and tested.", "No-go"],
        ["App Store submission completeness", "Privacy policy, account deletion, IAP metadata, screenshots, support URL, review notes, demo credentials and export/privacy declarations are complete.", "No-go"],
    ], widths=[1.55, 4.6, 0.85], font_size=8.4)

    doc.add_heading("P1 checks that should be completed before release", level=2)
    bullets(doc, [
        "Test the smallest and largest supported iPhones, dark/light appearance, landscape where supported and large Dynamic Type on scoring and purchase screens.",
        "Review all destructive flows: account deletion, match deletion/archive, user removal, subscription cancellation and Root Admin password/email operations.",
        "Confirm privacy disclosure matches actual analytics, notification polling, UserDefaults, account data, purchase identifiers and support email handling.",
        "Check network-loss behaviour around start/score/end, login/session expiry, purchase submission and notification reconciliation.",
        "Confirm enabled sports in production prevent access to placeholder engines through both UI and direct API calls.",
        "Prepare support replies for purchase pending, declined renewal, restore mismatch, subscription active but UI stale, and delayed App Store review.",
    ])

    doc.add_heading("Backend and database freeze policy", level=1)
    status_line(doc, "Freeze point", "End of Day 2, only after every backend/data P0 gate is green. Freezing earlier would preserve known release risk; freezing later would remove the stabilisation window.", NAVY)
    table(doc, ["Change class after freeze", "Policy", "Examples"], [
        ["Style-only", "Allowed after focused visual regression.", "Colour, spacing, typography, copy, non-functional imagery and layout that do not change accessibility semantics or control behaviour."],
        ["Client behaviour", "Default denied; release-owner exception required.", "Navigation, validation, session handling, StoreKit flow, offline queue, scoring actions or API timing."],
        ["Backend function/API", "Frozen; P0 incident fix only.", "Lambda code, routes, authorization, entitlement rules, notification processing, reconciliation or response contracts."],
        ["Database", "Frozen; emergency forward-only migration only.", "Tables, columns, constraints, triggers, indexes, RLS or production data correction."],
        ["Configuration", "Controlled as code/change record.", "Feature flags, CORS origins, secrets, Apple endpoints, alarms, timeouts and environment variables."],
    ], widths=[1.45, 2.15, 3.4])

    doc.add_heading("Exception procedure", level=2)
    bullets(doc, [
        "Record the user impact, severity and why the release cannot proceed safely without the change.",
        "Identify exact files, functions, schema objects, API contracts and production data affected.",
        "Provide a minimal patch, tests that fail before and pass after, deployment steps and rollback/forward-fix path.",
        "Obtain explicit release-owner and backend/data-owner approval before merge or migration.",
        "Re-run the complete release gate, not only the local test for the defect, and update the release manifest.",
    ])

    doc.add_heading("Database safety at freeze", level=2)
    bullets(doc, [
        "Record deployed migration versions and checksums; do not infer production state from the repository alone.",
        "Verify Supabase backups and point-in-time recovery retention, then perform a restore rehearsal into an isolated target.",
        "Capture row-count and integrity checks for subscriptions, entitlement audit, notification events, sessions, memberships, matches and action receipts.",
        "Prohibit manual production SQL except an approved incident runbook with before/after evidence and peer review.",
        "Prefer forward fixes for migrations; never rely on a destructive down migration for customer data recovery.",
    ])

    add_page(doc, "Root Admin security assessment")
    status_line(doc, "Current verdict", "Functionally useful and materially safer than the former trust-header design, but not safe enough for broad public exposure because it lacks strong second-factor and edge-level enforcement.", RED)
    doc.add_heading("Controls already present", level=2)
    bullets(doc, [
        "Opaque cryptographically random session tokens are stored as SHA-256 hashes rather than plaintext.",
        "Sessions have an absolute expiry, can be revoked, and a new login revokes the previous active session for that administrator.",
        "Root-admin routes call the shared session validator and the historical x-root-admin-request trust bypass has been removed.",
        "Passwords are checked with a recognised password-hashing implementation rather than stored or compared as plaintext.",
        "The browser currently keeps the root-admin token in sessionStorage, limiting persistence across a browser restart but not protecting it from script execution in the origin.",
    ])

    doc.add_heading("Material gaps", level=2)
    table(doc, ["Gap", "Risk", "Priority"], [
        ["No MFA", "One stolen/reused password can unlock platform-wide destructive and identity operations.", "Blocker"],
        ["No edge authorizer", "Every request reaches public Lambda routing before application code decides authorization; a missed handler check becomes critical.", "Blocker"],
        ["Token in sessionStorage", "An XSS defect can read and exfiltrate the bearer token.", "Blocker"],
        ["Eight-hour absolute session and no idle timeout", "A lost or unattended admin browser remains privileged for too long.", "Required"],
        ["No demonstrated login throttle/lockout", "Credential stuffing and password guessing are insufficiently contained.", "Blocker"],
        ["CORS allows every origin", "It is unnecessarily broad for a privileged portal and weakens browser isolation assumptions.", "Blocker"],
        ["No comprehensive admin mutation audit", "Password changes, verification, membership, sport, notification and delete actions lack one immutable accountability trail.", "Blocker"],
        ["No role separation", "Every administrator appears effectively all-powerful; compromise has maximum blast radius.", "Required"],
    ], widths=[1.55, 4.45, 1.0], font_size=8.5)

    doc.add_heading("Recommended production design", level=2)
    bullets(doc, [
        "Create a dedicated Amazon Cognito user pool for named HitnScore administrators only. Require TOTP MFA for every admin; do not permit self-registration.",
        "Attach an API Gateway HTTP API JWT authorizer to every /root_admin route and make authorization the route default. Keep Lambda role checks as defence in depth.",
        "For the browser, prefer a small backend-for-frontend that exchanges the identity response for a Secure, HttpOnly, SameSite cookie. Remove privileged bearer tokens from sessionStorage.",
        "Use a dedicated admin hostname, an explicit origin allowlist, strict Content Security Policy and standard security headers. Do not treat a secret or obscure URL as a control.",
        "Set a short privileged session: about one hour absolute and fifteen minutes idle, with re-authentication for password, email verification, membership deletion and destructive match operations.",
        "Rate-limit and alert on admin login and authorization failures. Apply account lockout/backoff without creating a trivial denial-of-service path.",
        "Write every privileged read of sensitive data and every mutation to an append-only audit stream with actor, target, before/after summary, request ID, IP/device context and outcome.",
        "Separate read-only support, billing, operations and security roles. Retain one break-glass account offline, with monitored use and a documented recovery process.",
    ])

    doc.add_heading("Root Admin plan for the next four days", level=1)
    table(doc, ["Time", "Minimum action", "Acceptance test"], [
        ["Day 1", "Choose Cognito + required TOTP as the release design. Create named admin identities and remove shared credentials. If this cannot be completed, decide that the portal/API remain network-restricted at launch.", "Two named admins can enroll TOTP; self-registration and password-only access are impossible."],
        ["Day 2", "Enforce the JWT authorizer on every privileged route; restrict CORS; add login throttling; shorten sessions; add audit events for every mutation; remove sessionStorage token handling through an HttpOnly cookie/BFF or equivalent protected flow.", "Anonymous, expired, wrong-audience, wrong-issuer and non-admin tokens are rejected before Lambda; every mutation produces one audit record."],
        ["Freeze", "Disable any unneeded high-risk admin route and freeze route/auth/audit configuration with the backend candidate.", "Automated inventory confirms every /root_admin route has the expected authorization and audit posture."],
        ["Day 3", "Pen-test the practical paths: direct API calls, privilege tampering, CSRF, XSS token access, brute-force handling, session expiry/revocation and destructive re-authentication.", "No password-only or browser-origin bypass; revoked and idle sessions fail; destructive events are attributable."],
        ["Day 4", "Enable alerts, export/retain audit records, document break-glass access and confirm network restriction if the durable design is incomplete.", "On-call receives a test alert and can revoke all admin sessions without a deployment."],
    ], widths=[0.85, 4.15, 2.0], font_size=8.2)
    status_line(doc, "Fallback rule", "If MFA plus API enforcement cannot be completed and tested by the Day 2 freeze, do not expose Root Admin publicly. Put both the portal and its API behind an approved VPN, identity-aware proxy or tightly controlled network path until the durable design is ready.", RED)

    add_page(doc, "Release candidate validation")
    doc.add_heading("Clean-machine automated gate", level=2)
    table(doc, ["Area", "Required command or outcome"], [
        ["Backend", "Install pinned test dependencies; run pytest -c testing/automated/backend/pytest.ini testing/automated/backend; compile all backend modules."],
        ["Frontend", "Install from the lockfile; run the production build and Playwright public-route smoke suite."],
        ["Native logic", "Run testing/automated/mobile/run-tennis-scenarios.sh and any checked-in squash/racketball unit scenarios."],
        ["Native release", "Create a signed Archive in Release configuration, validate it, upload to App Store Connect and install the resulting TestFlight build."],
        ["Infrastructure", "Build/validate SAM, produce a change set, review IAM/environment changes, deploy only from the tagged candidate and smoke every public and privileged route class."],
    ], widths=[1.3, 5.7])

    doc.add_heading("Manual release-candidate journey", level=2)
    bullets(doc, [
        "Register a new personal account, receive password setup, sign in, edit profile, reset password and delete the account through both confirmations.",
        "Sign in with multiple memberships, select/switch an association, verify tenant isolation and confirm an expired cached session cannot be restored.",
        "Create, start, score, undo and finish squash, racketball and representative tennis matches; verify compact/Dynamic Type layouts and completed summaries.",
        "Open an active match online, go offline, score and restart the app, reconnect, replay once and verify match_action_receipts prevent duplicate mutation.",
        "Confirm Personal Free sees only its latest three completed matches and cannot bypass the limit with a direct authenticated match read.",
        "Purchase Personal Plus, verify server activation, restore on a second device, reject a different HitnScore account, and execute the remaining exceptional lifecycle cases.",
        "Publish a notification to all users and to a plan; verify inbox polling, unread/read state and cross-device consistency.",
        "Exercise Root Admin with each production role, including denied actions, audit records, session expiry/revocation and the break-glass procedure.",
    ])

    doc.add_heading("App Store review package", level=2)
    bullets(doc, [
        "Provide working reviewer credentials and precise instructions for reaching paid content, scoring and account deletion.",
        "Explain any server-side feature switches or configuration needed for review; make the backend available for the review period.",
        "Ensure subscription title, duration, price, ongoing value, privacy policy and terms are visible and accurate.",
        "Confirm all digital features use in-app purchase and that restored entitlements work on every device for the signed-in customer.",
        "Upload final device screenshots, support URL, privacy answers, age rating and export-compliance answers from the exact release candidate.",
    ])

    add_page(doc, "Go / no-go decision record")
    table(doc, ["Question", "Required answer"], [
        ["Is the release commit/tag immutable and reproducible?", "Yes"],
        ["Are all P0 gates closed with evidence?", "Yes"],
        ["Is Root Admin MFA/edge enforcement complete, or is the surface network-restricted?", "Yes"],
        ["Did the signed TestFlight build pass the full journey against the production candidate?", "Yes"],
        ["Are production Apple configuration and remaining exceptional cases signed off?", "Yes"],
        ["Are backup, rollback, alarms, notification replay and support ownership rehearsed?", "Yes"],
        ["Are known exclusions reflected in product copy and reviewer notes?", "Yes"],
        ["Does any open defect risk data loss, unauthorized access, incorrect scoring, incorrect entitlement or inability to recover?", "No"],
    ], widths=[5.7, 1.3], font_size=9.2)
    status_line(doc, "Decision rule", "Any different answer is a no-go for public release. A controlled TestFlight beta may continue while the failing gate is resolved.", RED)

    doc.add_heading("Rollback and first 72 hours", level=2)
    table(doc, ["Window", "Monitor", "Action threshold"], [
        ["Before release", "Database backup/PITR, deployment versions, secret references, App Store notification endpoint health.", "Stop if restore or rollback cannot be demonstrated."],
        ["0–6 hours", "Login errors, 4xx/5xx, Lambda faults/timeouts, subscription verification/notification/reconciliation errors, scoring action failures.", "Pause/phased release or roll back on sustained abnormal errors, authorization failures or incorrect entitlements."],
        ["6–24 hours", "Purchase conversion, restore failures, unexpected plan changes, duplicate events, offline replay, support contacts.", "Disable purchasing if authoritative verification or entitlement state is uncertain; preserve existing paid access while investigating."],
        ["24–72 hours", "Crash reports, device/layout issues, admin security events, notification backlog, data integrity checks.", "Ship style-only remediation normally; use the freeze exception process for functional/backend changes."],
    ], widths=[1.0, 4.15, 1.85], font_size=8.5)

    doc.add_heading("Production incident principles", level=2)
    bullets(doc, [
        "Protect customer data and paid entitlement before protecting release dates.",
        "For Apple lifecycle uncertainty, reconcile with Apple before locally downgrading wherever the existing bounded logic supports it.",
        "Do not modify production rows manually to hide a symptom; preserve event evidence and use an approved replay or forward fix.",
        "Revoke privileged sessions and restrict the admin surface immediately if compromise is suspected.",
        "Keep the previous backend artifact and client release available; record the exact decision and customer impact for every rollback.",
    ])

    doc.add_heading("Prioritised backlog after launch", level=1)
    table(doc, ["Priority", "Item", "Why it remains"], [
        ["P1 · first week", "CI for backend, web and iOS archive/test", "The release gate must stop depending on one configured workstation."],
        ["P1 · first week", "Complete Root Admin RBAC, audit retention and periodic access review", "MFA is necessary but does not reduce blast radius or establish governance by itself."],
        ["P1 · first week", "Security scanning, dependency review, rate limits and broader API audit", "Public-route and privileged hardening remain incomplete."],
        ["P1 · first week", "Subscription reconciliation dashboards and replay tooling", "Operators need to identify and repair missed or failed lifecycle events safely."],
        ["P1 · first week", "iPhone scoring UX/device hardening", "The current layout is improved but not yet considered final across every device/accessibility setting."],
        ["P2", "APNs push and background badges", "Useful engagement/operations improvement, not required for the current inbox-polling launch."],
        ["P2", "Complete WebSocket live display", "Scaffolded capability should not be advertised until broadcast reliability is production-grade."],
        ["P2", "Finish settings scaffolds and shared profile photos", "Polish and feature breadth after the backend stabilisation period."],
    ], widths=[1.15, 3.0, 2.85], font_size=8.4)

    doc.add_heading("What may continue during the style-only period", level=2)
    bullets(doc, [
        "Brand colour, spacing, typography, icon alignment, non-functional imagery and copy improvements.",
        "Accessibility contrast and text sizing fixes that do not alter control meaning, data flow or scoring behaviour.",
        "App Store screenshots, marketing pages, help content and reviewer/support instructions.",
        "Visual regression fixes with before/after screenshots and a confirmation that API requests and persisted state are unchanged.",
    ])
    doc.add_paragraph(
        "Treat a change as functional—not stylistic—if it can alter which request is sent, when it is sent, what is stored, which plan or role is selected, whether a control is destructive, or how a score is calculated. Those changes remain frozen."
    )

    add_page(doc, "Evidence, assumptions and sources")
    doc.add_heading("Repository evidence", level=2)
    bullets(doc, [
        "AGENTS.md current product state and risk register, read 11 September 2026.",
        "docs/production-readiness-and-personal-plus-review-2026-09-08.md for the previous baseline.",
        "docs/apple-subscription-production.md and docs/app-store-subscription-production.md for production activation gates.",
        "backend/common/root_admin_session_logic.py, backend/functions/root_admin_login/handler.py, frontend/src/context/RootAdminContext.jsx and backend/template.yaml for the current privileged-session and API posture.",
        "Git history from 3–10 September 2026 and the 139-file change set between the root-admin security milestone and current HEAD.",
        f"Companion manual evidence report: {PRIOR_REPORT.name}.",
    ])

    doc.add_heading("Assumptions and caveats", level=2)
    bullets(doc, [
        "No direct production AWS, Supabase, App Store Connect or Apple signing configuration was changed or independently inspected for this assessment.",
        "Historical test results are labelled separately from commands executed on 11 September; missing local dependencies were not interpreted as source failures.",
        "The release date assumes the product owner accepts the narrow scope and can provide daily decisions, credentials, Apple agreements and an available signed-build host.",
        "Public availability depends on Apple review. Apple states that most submissions are reviewed quickly, but incomplete submissions can be delayed and the timing is not guaranteed.",
        "Security recommendations describe the target control outcome. The exact Cognito, cookie/BFF, domain and network design must be implemented and tested against the deployed topology before sign-off.",
    ])

    doc.add_heading("External guidance", level=2)
    links = [
        ("Apple App Review Guidelines", "https://developer.apple.com/app-store/review/guidelines/"),
        ("Apple App Review overview", "https://developer.apple.com/app-store/review/"),
        ("Apple submission process", "https://developer.apple.com/help/app-store-connect/manage-submissions-to-app-review/submit-an-app"),
        ("Apple release-build testing", "https://developer.apple.com/documentation/Xcode/testing-a-release-build"),
        ("Apple account deletion requirements", "https://developer.apple.com/support/offering-account-deletion-in-your-app"),
        ("AWS Cognito MFA", "https://docs.aws.amazon.com/cognito/latest/developerguide/user-pool-settings-mfa.html"),
        ("AWS API Gateway HTTP API JWT authorizers", "https://docs.aws.amazon.com/apigateway/latest/developerguide/http-api-jwt-authorizer.html"),
        ("AWS Lambda operational best practices", "https://docs.aws.amazon.com/lambda/latest/dg/best-practices.html"),
        ("OWASP Session Management Cheat Sheet", "https://cheatsheetseries.owasp.org/cheatsheets/Session_Management_Cheat_Sheet.html"),
        ("OWASP Authentication Cheat Sheet", "https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html"),
    ]
    for label, url in links:
        paragraph = doc.add_paragraph(style="List Bullet")
        add_hyperlink(paragraph, label, url)

    doc.add_heading("Final recommendation", level=1)
    doc.add_paragraph(
        "Use the next two days to remove the remaining backend, data, Apple and privileged-security uncertainty; freeze the resulting candidate at the end of Day 2; spend Day 3 proving that immutable candidate; and use Day 4 for controlled deployment, submission and operational rehearsal. If Root Admin cannot be protected with strong MFA plus API enforcement in that window, keep it off the public internet. If the subscription exception suite, signed archive or rollback evidence is incomplete, ship a controlled TestFlight beta rather than a paid public release."
    )

    headers_and_footers(doc)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
